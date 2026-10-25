import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime
from pathlib import Path
from typing import List, Dict, Any
import config


def _style_header(ws, col_count: int):
    header_fill = PatternFill(
        start_color="1F4E79", end_color="1F4E79", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def export_users_to_excel(users: List[Dict[str, Any]]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users Data"

    headers = [
        "User ID",
        "Username",
        "First Name",
        "Balance ($)",
        "Accounts Sold",
        "Total Withdrawn ($)",
        "Status",
        "Premium",
        "Joined At",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for u in users:
        status_str = "Banned" if u.get("is_banned") else "Active"
        prem_str = "Yes" if u.get("is_premium") else "No"
        ws.append(
            [
                u.get("user_id"),
                f"@{u.get('username')}" if u.get("username") else "N/A",
                u.get("first_name", ""),
                round(float(u.get("balance", 0.0)), 2),
                u.get("accounts_sold", 0),
                round(float(u.get("total_withdrawn", 0.0)), 2),
                status_str,
                prem_str,
                str(u.get("joined_at", "")),
            ]
        )

    _autofit_columns(ws)
    filename = (
        f"users_export_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    filepath = config.EXPORTS_DIR / filename
    wb.save(filepath)
    return str(filepath)


def export_withdrawals_to_excel(
    withdrawals: List[Dict[str, Any]], prefix: str = "withdrawals"
) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Withdrawals"

    headers = [
        "W-ID",
        "User ID",
        "Username",
        "Amount ($)",
        "Method",
        "Wallet / Account",
        "Status",
        "Admin Note",
        "Requested Date",
        "Processed Date",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for w in withdrawals:
        ws.append(
            [
                w.get("wid"),
                w.get("user_id"),
                f"@{w.get('username')}" if w.get("username") else "N/A",
                round(float(w.get("amount", 0.0)), 2),
                w.get("method", ""),
                w.get("wallet", ""),
                str(w.get("status", "")).upper(),
                w.get("admin_note", ""),
                str(w.get("created_at", "")),
                str(w.get("processed_at", "") or "Pending"),
            ]
        )

    _autofit_columns(ws)
    filename = f"{prefix}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = config.EXPORTS_DIR / filename
    wb.save(filepath)
    return str(filepath)


def export_accounts_to_excel(accounts: List[Dict[str, Any]]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts Report"

    headers = [
        "Account ID",
        "Phone Number",
        "User ID",
        "Username",
        "Country",
        "Category",
        "Payout ($)",
        "Status",
        "Downloaded",
        "Date Received",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for a in accounts:
        ws.append(
            [
                a.get("id"),
                a.get("phone"),
                a.get("user_id"),
                f"@{a.get('username')}" if a.get("username") else "N/A",
                a.get("country_name", "N/A"),
                a.get("category", ""),
                round(float(a.get("price", 0.0)), 2),
                a.get("status", ""),
                "Yes" if a.get("is_downloaded") else "No",
                str(a.get("created_at", "")),
            ]
        )

    _autofit_columns(ws)
    filename = (
        f"accounts_report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    filepath = config.EXPORTS_DIR / filename
    wb.save(filepath)
    return str(filepath)
